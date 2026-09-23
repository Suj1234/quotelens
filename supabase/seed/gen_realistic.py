# Realistic variants of the five responses + extra real-world samples. Prices identical to core (gold unchanged).
import json, os, random, math, subprocess, re
from gen_core import LINES, balaji, kohinoor, westline, orient, anand, QUESTIONS, QA, FX_USD
random.seed(99)
R="pack/realistic"; os.makedirs(R,exist_ok=True)
for d in ["01_balaji","02_kohinoor","03_westline","04_orientpack","05_anand","06_extra_samples"]: os.makedirs(f"{R}/{d}",exist_ok=True)
def inr(v): return f"{v:,.0f}".replace(",","X").replace("X",",")  # keep western grouping; Indian grouping below
def inr_in(v):
    s=f"{int(round(v))}"; 
    if len(s)<=3: return s
    head,tail=s[:-3],s[-3:]; parts=[]
    while len(head)>2: parts.insert(0,head[-2:]); head=head[:-2]
    if head: parts.insert(0,head)
    return ",".join(parts)+","+tail

# ---------------- V1 Balaji: messy Excel ----------------
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Sheet1"
ws["B1"]="SRI BALAJI PACKAGING"; ws["B1"].font=Font(name="Arial",bold=True,size=14,color="C00000")
ws["B2"]="Mfrs of Corrugated Boxes, Sheets & Partitions | An ISO 9001:2015 Co."; ws["B2"].font=Font(name="Arial",italic=True,size=9)
ws["B3"]="Plot No 42-B, SIPCOT Ind. Complex, Hosur - 635 126. Ph: 04344-27xxxx  email: sbp.hosur@gmail.com"
ws["B5"]="QUOTATION"; ws["B5"].font=Font(bold=True,underline="single")
ws["B6"]="Qtn No: SBP/26-27/0912"; ws["E6"]="Dt: 27.09.2026"
ws["B7"]="To, Meridian Foods Pvt Ltd, Bengaluru. Kind Attn: Mr Sujit Menon"
ws["B8"]="Ref: Your enquiry MER-0417 dtd 24.09.2026 - Corrugated Boxes"
ws["B9"]="Dear Sir, With ref to the above we are pleased to quote our lowest rates as under:"
hdr=["Sl","Our Code","Particulars","Size (LxBxH) mm","Ply","B.F.","Rate/1000 Nos (Rs)","GST","Remark"]
ws.append([]); ws.append(["",*hdr])
hr=ws.max_row
for c in ws[hr][1:]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="FFFF99"); c.border=Border(bottom=Side(style="thin"))
sl=0
for L in LINES:
    n=L["line_no"]; sl+=1
    size=f'{L["length_mm"]}x{L["width_mm"]}' + (f'x{L["height_mm"]}' if L["height_mm"] else "")
    part={"box":"Corrugated Box","sheet":"Corr. Sheet","partition":"Partition"}[L["item_type"]]
    if L["item_type"]=="partition": part=f'Partition {L["description"].split()[3]}'
    if "Layer pad" in L["description"]: part="Layer Pad"
    rate=balaji[n]
    # realistic clutter: some as text with /-; one "(revised)"; one Indian-grouped text; one number stored as text
    if n in (4,11): rate_cell=f"{inr_in(rate)}/-"
    elif n==7: rate_cell=f"{rate} (revised)"
    elif n==19: rate_cell=str(rate)  # number stored as text
    else: rate_cell=rate
    gst="18%" if n%9 else "18"
    rem={1:"Export qlty, 2 col print",6:"Heavy duty",12:"",16:"Plain",23:"Sheets in bundle of 50",28:"Loose"}.get(n,"")
    ws.append(["",sl,f"SBP-{1000+n}",part,size,f'{L["ply"]}',L["burst_factor"],rate_cell,gst,rem])
    if n==7: ws.cell(ws.max_row,8).comment=Comment("Revised from 13200 after paper cost increase - RK","RK")
    if n==15:  # stray blank + note row (realistic)
        ws.append([]); ws.append(["","","","NOTE: sizes are inside dimensions","","","","","",""]); ws.cell(ws.max_row,4).font=Font(italic=True,color="808080")
r=ws.max_row+2
for i,t in enumerate(["Terms & Conditions:","1) Rates are per 1000 Nos, basic, delivered at your Hosur / Nelamangala plant (freight incl.)","2) GST 18% extra",
 "3) Payment: 45 days from invoice dt as per your enquiry","4) Validity: 60 days","5) Spl. discount 3% on total bill value if all 30 items awarded to us for full year",
 "6) Die/tooling charges NIL for repeat sizes","7) MOQ 5000 nos per size per delivery","E.& O.E.","","Thanking you, For SRI BALAJI PACKAGING","","R. Krishnan (Manager - Marketing)  Mob 94433 xxxxx"]):
    ws.cell(r+i,2,t)
ws.cell(r,2).font=Font(bold=True,underline="single")
ws.column_dimensions["A"].width=3
for col,wd in zip("BCDEFGHIJ",[6,10,22,16,5,5,18,7,24]): ws.column_dimensions[col].width=wd
ws.row_dimensions[4].hidden=True
ws2=wb.create_sheet("Questionnaire"); ws2.append(["Q.No","Reply"])
for q,t,a,m,d in QUESTIONS: ws2.append([q,QA["balaji"][q][0]])
ws3=wb.create_sheet("Sheet3")  # empty sheet, very common
wb.save(f"{R}/01_balaji/Qtn SBP-0912 Meridian.xlsx")
open(f"{R}/01_balaji/balaji_cover_email.txt","w").write("""Dear Sir,

Kindly find attached our quotation for corrugated boxes against your enquiry MER-0417. Rates are our lowest and best. ISO certificate copy also attached for your ref.

Pls confirm receipt. Awaiting your valued order.

Thanks & Regards
R. Krishnan
Manager - Marketing
Sri Balaji Packaging, Hosur
Mob: 94433 xxxxx

Disclaimer: This e-mail and any files transmitted with it are confidential and intended solely for the use of the individual or entity to whom they are addressed.
""")

# ---------------- V2 Kohinoor: PDF with logo, stamp, signature, strikethrough correction ----------------
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas as rlcanvas
from reportlab.lib import colors
from reportlab.lib.units import mm
def kohinoor_real():
    path=f"{R}/02_kohinoor/KC_Quotation_1187_Meridian.pdf"; c=rlcanvas.Canvas(path,pagesize=A4); W,H=A4
    def head():
        c.setFillColor(colors.HexColor("#1F3A5F")); c.rect(15*mm,H-30*mm,22*mm,18*mm,fill=1,stroke=0)
        c.setFillColor(colors.white); c.setFont("Helvetica-Bold",20); c.drawCentredString(26*mm,H-23*mm,"KC")
        c.setFillColor(colors.HexColor("#1F3A5F")); c.setFont("Times-Bold",17); c.drawString(41*mm,H-18*mm,"Kohinoor Corrugators Pvt. Ltd.")
        c.setFillColor(colors.black); c.setFont("Times-Roman",8.5)
        c.drawString(41*mm,H-23*mm,"Gat No. 118, Chakan MIDC Phase-II, Tal. Khed, Dist. Pune - 410 501 (MS)   Tel: 02135-65xxxx   E-mail: sales@kohinoorcorr.co.in")
        c.drawString(41*mm,H-27*mm,"GSTIN: 27AABCK2210H1ZP     CIN: U21012PN1994PTC081xxx     An ISO 9001:2015 & BRCGS Certified Company")
        c.setStrokeColor(colors.HexColor("#1F3A5F")); c.setLineWidth(1.2); c.line(15*mm,H-32*mm,W-15*mm,H-32*mm)
    head()
    c.setFont("Times-Bold",11); c.drawString(15*mm,H-40*mm,"QUOTATION"); c.setFont("Times-Roman",9.5)
    c.drawString(15*mm,H-46*mm,"Ref. No.: KC/Q/26-27/1187"); c.drawString(120*mm,H-46*mm,"Date: 29/09/2026")
    c.drawString(15*mm,H-52*mm,"To,"); c.drawString(15*mm,H-56.5*mm,"M/s. Meridian Foods Pvt. Ltd., Bengaluru."); c.drawString(15*mm,H-61*mm,"Kind Attn.: Mr. Sujit Menon (Purchase)")
    c.drawString(15*mm,H-67*mm,"Sub: Quotation for Corrugated Boxes & Sheets against your RFx MER-0417 dtd. 24.09.2026")
    c.drawString(15*mm,H-73*mm,"Dear Sir,  With reference to the above, we are pleased to submit our most competitive rates as under. Sr. No. 28, 29 & 30")
    c.drawString(15*mm,H-77.5*mm,"(partition sets) are not in our manufacturing range, hence not quoted.")
    # table
    y=H-84*mm; xs=[15,25,110,150,168,195]
    c.setFont("Times-Bold",9); c.setFillColor(colors.HexColor("#E8EEF7")); c.rect(15*mm,y-5*mm,180*mm,6*mm,fill=1,stroke=0); c.setFillColor(colors.black)
    for x,t in zip(xs,["Sr.","Description of Item","Size (mm)","Ply/BF","Rate Rs./1000 Nos *"]): c.drawString(x*mm+1*mm,y-3.5*mm,t)
    c.setFont("Times-Roman",8.6); rowh=5.6*mm; y-=5*mm
    for L in LINES:
        n=L["line_no"]
        if n not in kohinoor: continue
        y-=rowh; size=f'{L["length_mm"]} x {L["width_mm"]}' + (f' x {L["height_mm"]}' if L["height_mm"] else "")
        c.drawString(xs[0]*mm+1*mm,y+1.5*mm,str(n)); c.drawString(xs[1]*mm+1*mm,y+1.5*mm,L["description"]); c.drawString(xs[2]*mm+1*mm,y+1.5*mm,size)
        c.drawString(xs[3]*mm+1*mm,y+1.5*mm,f'{L["ply"]}/{L["burst_factor"]}')
        val=inr_in(kohinoor[n])
        if n==17:  # strikethrough old value, corrected value written next to it
            old=inr_in(kohinoor[n]+120); c.drawRightString(178*mm,y+1.5*mm,old); tw=c.stringWidth(old,"Times-Roman",8.6)
            c.setStrokeColor(colors.black); c.setLineWidth(0.6); c.line(178*mm-tw,y+2.6*mm,178*mm,y+2.6*mm)
            c.setFont("Helvetica-Bold",8.6); c.setFillColor(colors.HexColor("#1a3a9a")); c.drawRightString(194*mm,y+1.5*mm,val); c.setFillColor(colors.black); c.setFont("Times-Roman",8.6)
        else: c.drawRightString(194*mm,y+1.5*mm,val)
        c.setStrokeColor(colors.grey); c.setLineWidth(0.3); c.line(15*mm,y,195*mm,y)
    for x in xs+[195]: c.line(x*mm,H-89*mm,x*mm,y)
    c.line(15*mm,H-89*mm,195*mm,H-89*mm)
    c.setFont("Times-Italic",7.8); c.drawString(15*mm,y-5*mm,"* Rates are net of our 2.5% early-payment discount, applicable only if payment is received within 10 days of invoice. For payment beyond 10 days,")
    c.drawString(15*mm,y-8.5*mm,"  rates are to be read as quoted / 0.975.   ** Sr. 17 rate corrected by hand - please read Rs. "+inr_in(kohinoor[17])+".   E.&O.E.")
    c.setFont("Times-Roman",7); c.drawRightString(W-15*mm,10*mm,"Page 1 of 2"); c.showPage()
    head(); c.setFont("Times-Bold",10.5); c.drawString(15*mm,H-40*mm,"Terms & Conditions:"); c.setFont("Times-Roman",9.5); y=H-46*mm
    for t in ["1. Rates: Rs. per 1000 Nos., delivered at your Hosur plant. Delivery to Nelamangala plant through our Bengaluru godown at same rate.",
              "2. Taxes: GST @ 18% extra as applicable.","3. Validity: 30 days from date of quotation.",
              "4. Payment: 30 days from date of invoice. 45 days acceptable, however early payment discount will not be applicable (refer note on page 1).",
              "5. Minimum order: 10,000 Nos. per size per delivery. Delivery: 10 days from receipt of confirmed PO / art work approval.",
              "6. Tolerance: +/- 5% in quantity supplied. Sizes are inside dimensions.","7. Force majeure conditions apply."]:
        c.drawString(15*mm,y,t); y-=5.5*mm
    y-=4*mm; c.setFont("Times-Bold",10.5); c.drawString(15*mm,y,"Reply to Supplier Questionnaire:"); c.setFont("Times-Roman",9); y-=6*mm
    for q,t,a,m,d in QUESTIONS:
        c.drawString(15*mm,y,f"Q{q}. {t}"); y-=4.5*mm; c.setFont("Times-Bold",9); c.drawString(22*mm,y,f"Ans: {QA['kohinoor'][q][0]}"); c.setFont("Times-Roman",9); y-=6*mm
    y-=4*mm; c.drawString(15*mm,y,"We trust our rates are competitive and look forward to your valued order."); y-=8*mm
    c.drawString(15*mm,y,"Thanking you,"); y-=5*mm; c.drawString(15*mm,y,"For KOHINOOR CORRUGATORS PVT. LTD.")
    # stamp (circle, blue) + signature scribble
    c.saveState(); c.translate(60*mm,y-16*mm); c.rotate(-12); c.setStrokeColor(colors.HexColor("#2b4fb3")); c.setLineWidth(1.1); c.circle(0,0,13*mm); c.circle(0,0,10.5*mm)
    c.setFillColor(colors.HexColor("#2b4fb3")); c.setFont("Helvetica-Bold",6.5); c.drawCentredString(0,-1*mm,"PUNE"); c.setFont("Helvetica",5.2)
    import math as _m
    txt="* KOHINOOR CORRUGATORS PVT LTD *"; 
    for i,ch in enumerate(txt):
        ang=180-(i*360/len(txt)); c.saveState(); c.rotate(ang); c.drawCentredString(0,11.7*mm,ch); c.restoreState()
    c.restoreState()
    c.saveState(); c.setStrokeColor(colors.HexColor("#1a2f8a")); c.setLineWidth(1.3); p=c.beginPath(); x0,y0=100*mm,y-14*mm; p.moveTo(x0,y0)
    for i in range(1,26): p.curveTo(x0+i*1.6*mm,y0+random.uniform(-4,6)*mm,x0+i*1.9*mm,y0+random.uniform(-5,5)*mm,x0+i*2.2*mm,y0+random.uniform(-2,3)*mm)
    c.drawPath(p,stroke=1,fill=0); c.restoreState()
    c.setFont("Times-Roman",9); c.drawString(100*mm,y-20*mm,"R. Deshpande"); c.drawString(100*mm,y-24*mm,"Manager - Sales & Marketing"); c.drawString(100*mm,y-28*mm,"Mob: 98220 xxxxx")
    c.setFont("Times-Roman",7); c.drawRightString(W-15*mm,10*mm,"Page 2 of 2"); c.save()
kohinoor_real()

# ---------------- V3 Westline: docx with mixed fonts, chatty, /- numbers, pasted Excel-style table ----------------
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
d=Document()
p=d.add_paragraph(); r=p.add_run("WESTLINE PACKAGING"); r.bold=True; r.font.size=Pt(16); r.font.name="Arial"; r.font.color.rgb=RGBColor(0x2E,0x7D,0x32)
p=d.add_paragraph(); r=p.add_run("Plot 17, Changodar GIDC, Sarkhej-Bavla Highway, Ahmedabad 382213 | GSTIN 24AADCW7781R1Z2 | westlinepack@yahoo.co.in"); r.font.size=Pt(8.5); r.font.name="Calibri"
d.add_paragraph("Ref: WP/OFR/26-27/0344                                                                Date: 30-09-2026")
d.add_paragraph("To,\nMeridian Foods Pvt Ltd\nBengaluru\nAttn: Mr Sujit Menon ji")
p=d.add_paragraph(); r=p.add_run("Sub: Our offer for corrugated boxes, sheets & partitions – your RFx MER-0417"); r.bold=True; r.font.name="Arial"
d.add_paragraph("Respected Sir,")
d.add_paragraph("Thanks for your enquiry. Pls find below our best offer. All prices are per bundle (as packed, strapped & palletised) in INR, GST 18% extra, valid 60 days. Payment 45 days from invoice is ok for us. Freight to Hosur and Nelamangala both is included in the below bundle prices, delivery by road, our own transporter.")
def ph(n):
    L=LINES[n-1]; wl=westline[n]; short=L["description"].split("-ply ")[1]
    return f'item {n} ({short}) Rs {inr_in(wl["per_bundle"])}/- per bundle of {wl["bundle_size"]} nos' if wl["stated"] else None
p=d.add_paragraph(); r=p.add_run("5 PLY BOXES (your items 1 to 12): "); r.bold=True
p.add_run("; ".join([ph(n) for n in range(1,13) if ph(n)]) + f'. Item 5 and item 9 we can offer at Rs {inr_in(westline[5]["per_bundle"])}/- and Rs {inr_in(westline[9]["per_bundle"])}/- per bundle resp.')
p=d.add_paragraph(); r=p.add_run("3 PLY BOXES (items 13 to 22): "); r.bold=True
p.add_run("; ".join([f'item {n} Rs {inr_in(westline[n]["per_bundle"])}/- per bundle of {westline[n]["bundle_size"]}' for n in range(13,23) if westline[n]["stated"]]) + f'. For item 15 & 19 our price is Rs {inr_in(westline[15]["per_bundle"])}/- and Rs {inr_in(westline[19]["per_bundle"])}/- per bundle.')
d.add_paragraph("Sheets / layer pad / partitions as per below table (copied from our system):")
tbl=d.add_table(rows=1,cols=5); tbl.style="Light Grid Accent 1"
for i,t in enumerate(["Item","Product","Pack","Rate/bundle","Remarks"]): tbl.rows[0].cells[i].text=t
for n in range(23,31):
    L=LINES[n-1]; c=tbl.add_row().cells; c[0].text=str(n); c[1].text=L["description"]; c[2].text=f'{westline[n]["bundle_size"]} nos'; c[3].text=f'{inr_in(westline[n]["per_bundle"])}/-'; c[4].text=("Loose supply also possible" if n>=28 else "")
for row in tbl.rows:
    for cell in row.cells:
        for pp in cell.paragraphs:
            for rr in pp.runs: rr.font.size=Pt(9); rr.font.name="Calibri"
d.add_paragraph("")
p=d.add_paragraph(); r=p.add_run("Your questionnaire – our answers:"); r.bold=True; r.font.name="Arial"
for q,t,a,m,dq in QUESTIONS:
    ans=QA["westline"][q][0]
    if q==6: ans="No. BRC audit is planned for Q1 2027, we will share certificate once received."
    d.add_paragraph(f"{q}) {ans}")
d.add_paragraph("Hope the above is in line with your requirement. Kindly do the needful and revert with your valued order. For any clarification pls call undersigned.")
d.add_paragraph("Thanks & Regards,\nHemant Shah\nDirector – Sales\nWestline Packaging, Ahmedabad\nM: 98250 xxxxx")
d.save(f"{R}/03_westline/Westline offer Meridian Foods Sept26.docx")
open(f"{R}/03_westline/westline_clarification_reply.txt","w").write(open("pack/03_westline/westline_clarification_reply.txt").read().replace("Dear Mr. Menon,\n\nApologies for the omission.","Sujit ji,\n\nSorry for the confusion, our mistake.").replace("Regards,","Rgds,"))

# ---------------- V4 OrientPack: second synthetic photo (rotated portrait, different light) + pen note instructions ----------------
from PIL import Image, ImageOps, ImageFilter, ImageEnhance, ImageDraw
im=Image.open("pack/04_orientpack/OrientPack_Rate_Card_PHOTO_synthetic.jpg")
im2=im.rotate(90,expand=True); im2=ImageEnhance.Brightness(im2).enhance(0.85).filter(ImageFilter.GaussianBlur(0.9))
im2.save(f"{R}/04_orientpack/IMG_20261001_114532.jpg",quality=78)
open(f"{R}/04_orientpack/PEN_NOTE_INSTRUCTIONS.txt","w").write("""Before photographing the printed rate card (OrientPack_Rate_Card_PRINT_ME.pdf), write in blue/black pen, by hand, in the margin near the bottom:

   "FOB Chennai - freight extra. Valid 60 days. - R.K. 1/10"

and put a small tick mark next to 3-4 random rows. Rest your thumb over the line 14 price when you shoot.
This gives the extractor real handwriting (a term that is NOT in the printed text) and a real obstruction.
Expected: extraction picks up 'freight extra' from the handwritten note into response_terms.freight_terms_raw (bonus, not required by gold).
""")

# ---------------- V5 Anand: Hinglish email with quoted thread + mobile signature ----------------
rfx_mail=open("pack/00_rfx/rfx_meta.json").read()
anand_real=f"""Sujit sir good morning,

ref ur mail MER-0417. our rates as below

5 ply boxes (item 1 to 12) - Rs.42/- per kg
3 ply boxes (item 13 to 22) - Rs 38/- per kg

sheets, layer pad and partition (item 23 to 30) rest same as last year only, no change from our side.

freight extra at actual. GST 18% extra. payment 45 days ok. rate valid 60 days.

questionaire -
1 ISO 9001 yes
2 yes own corrugator
3 aprox 250 ton per month
4 one week for sample
5 local bakeries, Sri Krishna Sweets, Adyar Ananda Bhavan
6 BRC audit is in process, expecting by Dec
8 10-12 days
9 only Nelamangala regular, Hosur on request
10 yes

pls confirm and send PO. thanks

Anand Kumar
Anand Box Works, Peenya
98450 xxxxx

Sent from my iPhone

> On 24-Sep-2026, at 10:14 AM, Sujit Menon <sender+rfx-mer-0417-anand@gmail.com> wrote:
>
> Dear Anand Kumar,
>
> Meridian Foods invites quotations for its FY26-27 corrugated packaging requirement: 30 SKUs of 3-ply and 5-ply shipper, export, retail and inner cartons, corrugated sheets, layer pads and partition sets, for delivery to our Hosur (TN) and Nelamangala (KA) plants.
>
> Commercial terms: INR, per 1000 pieces, delivered to plant with freight included, payment 45 days, validity 60 days, 12-month contract. Response deadline: 07 October 2026.
>
> The line-item sheet (Excel) and supplier questionnaire (PDF) are attached. Please reply to this email with your quotation in any format convenient to you - we will process it as sent.
>
> Regards,
> Sujit Menon | Category Buyer - Packaging | Meridian Foods Pvt Ltd
"""
open(f"{R}/05_anand/anand_email_realistic.txt","w").write(anand_real)

# ---------------- 06 extra samples ----------------
X=f"{R}/06_extra_samples"
# (a) WhatsApp screenshot quote from a 6th vendor for 8 lines
def whatsapp_png():
    W,H=1080,1900; img=Image.new("RGB",(W,H),(11,20,26)); dr=ImageDraw.Draw(img)
    from PIL import ImageFont
    def font(sz,bold=False):
        for p in ["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]:
            if os.path.exists(p): return ImageFont.truetype(p,sz)
        return ImageFont.load_default()
    dr.rectangle([0,0,W,150],fill=(31,44,52)); dr.ellipse([40,40,110,110],fill=(120,140,150)); dr.text((130,45),"Sunrise Packers - Ravi",font=font(34,True),fill="white"); dr.text((130,95),"online",font=font(24),fill=(180,190,195))
    y=190
    def bubble(text,right,y,ts):
        f=font(30); lines=[]
        for para in text.split("\n"):
            cur=""
            for w in para.split(" "):
                if dr.textlength(cur+" "+w,font=f)>620: lines.append(cur); cur=w
                else: cur=(cur+" "+w).strip()
            lines.append(cur)
        h=len(lines)*40+60; x0=W-720 if right else 60; col=(0,92,75) if right else (32,44,51)
        dr.rounded_rectangle([x0,y,x0+660,y+h],radius=22,fill=col)
        for i,l in enumerate(lines): dr.text((x0+22,y+18+i*40),l,font=f,fill=(232,236,238))
        dr.text((x0+660-110,y+h-34),ts,font=font(20),fill=(170,180,185)); return y+h+22
    y=bubble("Ravi ji, Meridian Foods here. Can you quote for the 3 ply retail cartons? Sizes mailed to you. Need rate per 1000 nos delivered Hosur.",True,y,"10:02")
    y=bubble("Sure sir. Give me 1 hr",False,y,"10:05")
    lines8=[16,17,21,13,14,22,19,15]
    txt="Rates for Meridian (3 ply, BF 22/25, delivered Hosur):\n"
    for n in lines8:
        L=LINES[n-1]; rate=round(L["base_price_per_1000"]*0.96/10)*10
        txt+=f"{L['length_mm']}x{L['width_mm']}x{L['height_mm']} - Rs {inr_in(rate)}/1000\n"
    txt+="GST extra. 30 days payment. Valid 15 days only, paper rate is going up, pls understand"
    y=bubble(txt,False,y,"11:14")
    y=bubble("Item 15 (350x250x180) rate seems high? Kohinoor is lower",True,y,"11:20")
    L=LINES[14]; y=bubble(f"ok for 350x250x180 we can do Rs {inr_in(round(L['base_price_per_1000']*0.93/10)*10)}/1000 final. Others cannot reduce.",False,y,"11:26")
    y=bubble("Noted, will revert",True,y,"11:27")
    dr.rectangle([0,H-120,W,H],fill=(31,44,52)); dr.rounded_rectangle([40,H-95,W-160,H-25],radius=35,fill=(42,57,66)); dr.text((70,H-78),"Type a message",font=font(28),fill=(130,140,145))
    img.save(f"{X}/whatsapp_quote_sunrise_packers.png")
    return lines8
wa_lines=whatsapp_png()

# (b) scanned image-only PDF of the Kohinoor quotation page 1 (rotated, noisy)
subprocess.run(["pdftoppm","-r","110","-png","-f","1","-l","1",f"{R}/02_kohinoor/KC_Quotation_1187_Meridian.pdf",f"{X}/_scan"],check=True)
sp=[f for f in os.listdir(X) if f.startswith("_scan")][0]
sc=Image.open(f"{X}/{sp}").convert("L").rotate(1.6,expand=False,fillcolor=255)
import numpy as np
arr=np.array(sc).astype(np.int16); arr=arr-8+np.random.randint(-14,15,arr.shape,dtype=np.int16); sc=Image.fromarray(np.clip(arr,0,255).astype("uint8"))
sc=sc.filter(ImageFilter.GaussianBlur(0.7)); sc=ImageEnhance.Contrast(sc).enhance(1.15)
sc.convert("RGB").save(f"{X}/Kohinoor_quotation_SCANNED.pdf","PDF",resolution=110.0); os.remove(f"{X}/{sp}")

# (c) wrong-category quote (IT hardware) — should land in Unmatched
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Quote"
ws.append(["TechServe Solutions Pvt Ltd - Quotation Q-2026-0788 - Meridian Foods IT refresh"]); ws.append([])
ws.append(["S.No","Item","Specification","Qty","Unit price (INR)","GST","Total"])
for i,(it,sp_,q,pr) in enumerate([("Laptop","Dell Latitude 5450, i5, 16GB, 512GB",25,78500),("Docking station","Dell WD22TB4",25,18200),("Monitor 24 inch","Dell P2424HE",25,15900),("Wireless mouse","Logitech M331",40,1150),("Laptop bag","Targus 15.6",25,1650)],1):
    ws.append([i,it,sp_,q,pr,"18%",q*pr])
ws.append([]); ws.append(["Delivery 2 weeks. Payment 30 days. Validity 15 days. Warranty 3 yrs onsite."])
wb.save(f"{X}/wrong_category_IT_quote.xlsx")

# (d) revised offer email from Balaji (3 prices change) — tests conflict/supersede
rev={7:balaji[7]-400, 16:balaji[16]-150, 23:balaji[23]+300}
open(f"{X}/balaji_revised_offer_email.txt","w").write(f"""Dear Sir,

Further to our qtn SBP/26-27/0912, pls note revised rates for below items only (all other items unchanged):

Item 7 (Tray carton 600x400x120, 5 ply): Rs {inr_in(rev[7])}/- per 1000 nos (earlier {inr_in(balaji[7])})
Item 16 (Retail carton 200x150x100, 3 ply): Rs {inr_in(rev[16])}/- per 1000 nos
Item 23 (5 ply sheet 1200x800): Rs {inr_in(rev[23])}/- per 1000 nos - increase due to kraft paper price

Other terms same as our quotation. Kindly consider.

Thanks & Regards
R. Krishnan, Sri Balaji Packaging
""")

# (e) brochure — not a quote
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
ss=getSampleStyleSheet(); doc=SimpleDocTemplate(f"{X}/OrientPack_Product_Brochure.pdf",pagesize=A4)
doc.build([Paragraph("OrientPack Ltd — Product Range",ss["Title"]),Spacer(1,12),
 Paragraph("Regular slotted containers, die-cut cartons, heavy-duty 7-ply pallet boxes, printed shelf-ready packaging, corrugated sheets and partitions. Plants at Chennai (India) and Shah Alam (Malaysia). Capacity 1,400 MT/month.",ss["Normal"]),Spacer(1,12),
 Paragraph("Certifications: ISO 9001:2015, ISO 14001, BRCGS Packaging A+. Contact: sales.in@orientpack.example",ss["Normal"]),Spacer(1,12),
 Paragraph("This brochure does not constitute an offer. Prices on request.",ss["Italic"])])

# (f) partial responder: 6th vendor quotes only 5 lines by email, in per-piece paise style
open(f"{X}/sunrise_packers_email_partial.txt","w").write(f"""Hello Sujit,

As discussed on whatsapp, sharing rates for the retail cartons only. We dont make 5 ply.

Item 16 200x150x100 - Rs {LINES[15]['base_price_per_1000']*0.96/1000:.2f} per pc
Item 17 180x120x90 - Rs {LINES[16]['base_price_per_1000']*0.96/1000:.2f} per pc
Item 21 220x160x110 - Rs {LINES[20]['base_price_per_1000']*0.96/1000:.2f} per pc
Item 13 300x200x150 - Rs {LINES[12]['base_price_per_1000']*0.96/1000:.2f} per pc
Item 14 250x180x120 - Rs {LINES[13]['base_price_per_1000']*0.96/1000:.2f} per pc

Ex works Peenya. Transport extra approx Rs 1200 per trip to Hosur. GST extra. Advance 50% for first order.
ISO - applied, not yet received. BRC - no.

Ravi
Sunrise Packers, Bengaluru
""")

# expected-behaviour notes
json.dump({
 "realistic_five":{"note":"Prices identical to core; gold/gold.json applies unchanged. Differences are presentation only: clutter, text-formatted numbers, notes, stamps, strikethrough, Hinglish.",
   "balaji":{"file":"01_balaji/Qtn SBP-0912 Meridian.xlsx","traps":["rows 4 hidden","rates for items 4 & 11 stored as text '13,710/-' (Indian grouping)","item 7 rate cell text '<n> (revised)' with a cell comment","item 19 number stored as text","stray NOTE row after item 15","empty Sheet3","GST column mixes '18%' and 18"],"expected":"all 30 prices parse to the same numbers as gold; 'revised' text must not change the value"},
   "kohinoor":{"file":"02_kohinoor/KC_Quotation_1187_Meridian.pdf","traps":["logo block, stamp, signature scribble (vision noise)","item 17 old value struck through, corrected value printed beside it in blue + footnote **","footnote about net-of-discount rates","Times font, Indian date 29/09/2026"],"expected":"item 17 = corrected value (gold); footnote discount_treatment as in gold"},
   "westline":{"file":"03_westline/Westline offer Meridian Foods Sept26.docx","traps":["'Rs 1,709/-' number style","'per bundle of 25 nos'","bundle size missing for items 5, 9, 15, 19","Hinglish phrasing","table style Light Grid"],"expected":"same as gold"},
   "orientpack":{"file":"04_orientpack/IMG_20261001_114532.jpg","traps":["rotated 90 degrees (portrait photo)","darker, blurrier"],"expected":"auto-orient then same as gold; line 14 low_confidence"},
   "anand":{"file":"05_anand/anand_email_realistic.txt","traps":["'Rs.42/- per kg'","'aprox', 'questionaire' typos","quoted RFx thread below must be stripped (its terms are the BUYER's, not the vendor's)","'Sent from my iPhone'"],"expected":"same as gold; quoted block must not be extracted as vendor terms"}},
 "extra_samples":{
   "whatsapp_quote_sunrise_packers.png":{"vendor":"new vendor (not invited) — create on assign","expected":"classify quotation; 8 items extracted from chat bubbles with sizes only; item 15 has two values (first, then 'final' reduced) → keep final, flag conflict or note; unit per 1000 INR; validity 15 days → validity_short; payment 30 days; lines map by size","lines":wa_lines},
   "Kohinoor_quotation_SCANNED.pdf":{"expected":"image-only PDF (no text layer) → Gemini vision path; same 27 prices as gold within tolerance; skew must not break it"},
   "wrong_category_IT_quote.xlsx":{"expected":"classify quotation (it has prices) but every item → none_of_these → Unmatched items; grid unchanged; no crash"},
   "balaji_revised_offer_email.txt":{"expected":"second response from same vendor; items 7, 16, 23 → conflict review items (old vs revised); on confirm, revised values replace; other 27 untouched","revised":rev},
   "OrientPack_Product_Brochure.pdf":{"expected":"classify supporting or not_relevant; contains_prices=false; no extraction"},
   "sunrise_packers_email_partial.txt":{"expected":"per-piece prices ×1000; 5 of 30 lines; ex-works + transport per trip (cannot be converted to per-1000 without trip size → freight_treatment review); ISO 'applied' → Q1 ambiguous/no; BRC no → disqualified"}
 }},open(f"{R}/EXPECTED_BEHAVIOUR.json","w"),indent=1)
print("ok")
